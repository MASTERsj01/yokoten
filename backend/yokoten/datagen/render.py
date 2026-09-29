"""Render the document IR to PDF / DOCX / XLSX / MD and scanned PNG / JPG images."""

import html
import io
import math
import random
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from yokoten.datagen.world import COMPANY

FIXED_DATE = datetime(2025, 1, 1)
DOC_TYPE_LABEL = {
    "8d": "8D Problem-Solving Report",
    "lessons_learned": "Lessons Learned Report",
    "field_failure": "Field Failure Analysis",
    "test_report": "Test Report",
    "dfmea": "Design FMEA",
    "pfmea": "Process FMEA",
    "dvpr": "DVP&R",
    "design_review": "Design Review Minutes",
    "ecn": "Engineering Change Notice",
    "supplier_quality": "Supplier Quality Report",
    "work_instruction": "Work Instruction",
    "drawing": "Engineering Drawing",
    "inspection_record": "Incoming Inspection Record",
}


@dataclass
class Doc:
    id: str
    title: str
    doc_type: str
    fmt: str
    year: int
    date: str
    product_line: str | None = None
    component: str | None = None
    project: str | None = None
    plant: str | None = None
    suppliers: list[str] = field(default_factory=list)
    part_numbers: list[str] = field(default_factory=list)
    revision: str = "A"
    supersedes: str | None = None
    superseded_by: str | None = None
    classification: str = "internal"
    author: str = ""
    blocks: list = field(
        default_factory=list
    )  # ("h", lvl, text) ("p", text) ("list", [..]) ("kv", [(k, v)]) ("table", header, rows, caption) ("figure", png, caption)
    sheets: list = field(default_factory=list)  # xlsx: [(name, header, rows)]
    scan: dict = field(default_factory=dict)  # png/jpg: {"kind": ..., ...}

    @property
    def filename(self) -> str:
        return f"{self.id}_Rev{self.revision}.{self.fmt}"

    def meta(self) -> dict:
        return {
            "doc_id": self.id,
            "title": self.title,
            "doc_type": self.doc_type,
            "format": self.fmt,
            "file": f"{self.doc_type}/{self.filename}",
            "year": self.year,
            "date": self.date,
            "product_line": self.product_line,
            "component": self.component,
            "project": self.project,
            "plant": self.plant,
            "suppliers": self.suppliers,
            "part_numbers": self.part_numbers,
            "revision": self.revision,
            "supersedes": self.supersedes,
            "superseded_by": self.superseded_by,
            "classification": self.classification,
            "author": self.author,
        }


def render(doc: Doc, out_dir: Path, rng: random.Random) -> Path:
    path = out_dir / doc.doc_type / doc.filename
    path.parent.mkdir(parents=True, exist_ok=True)
    {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "md": _md, "png": _scan, "jpg": _scan}[doc.fmt](
        doc, path, rng
    )
    return path


def _png(fig) -> bytes:
    return fig() if callable(fig) else fig


def _header_kv(doc: Doc) -> list[tuple[str, str]]:
    return [
        ("Document", doc.id),
        ("Revision", doc.revision),
        ("Date", doc.date),
        ("Classification", doc.classification.upper()),
    ] + ([("Supersedes", doc.supersedes)] if doc.supersedes else [])


# ---------------------------------------------------------------- PDF (PyMuPDF Story: HTML -> PDF)
CSS = """
body { font-family: sans-serif; font-size: 10pt; line-height: 1.35; }
h1 { font-size: 16pt; margin: 4pt 0 8pt 0; }
h2 { font-size: 12.5pt; margin: 10pt 0 4pt 0; }
h3 { font-size: 11pt; margin: 8pt 0 3pt 0; }
p { margin: 0 0 6pt 0; }
table { border-collapse: collapse; margin: 4pt 0 8pt 0; }
td, th { border: 0.6pt solid #777; padding: 2pt 4pt; font-size: 9pt; vertical-align: top; }
th { background-color: #e6ebf1; text-align: left; }
.band { font-size: 8.5pt; color: #3a4a5c; border-bottom: 1pt solid #3a4a5c; padding-bottom: 3pt; margin-bottom: 6pt; }
.cap { font-size: 8.5pt; font-style: italic; margin: 2pt 0 8pt 0; }
"""


def _esc(s) -> str:
    return html.escape(str(s))


def _table_html(header, rows) -> str:
    out = "<table><tr>" + "".join(f"<th>{_esc(h)}</th>" for h in header) + "</tr>"
    for r in rows:
        out += "<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in r) + "</tr>"
    return out + "</table>"


def _pdf(doc: Doc, path: Path, rng: random.Random) -> None:
    import pymupdf

    archive = pymupdf.Archive()
    parts = [
        f'<div class="band"><b>{_esc(COMPANY)}</b> &#8212; {_esc(DOC_TYPE_LABEL[doc.doc_type])}'
        f" &#183; {_esc(doc.id)} &#183; Rev {_esc(doc.revision)} &#183; {_esc(doc.classification.upper())}</div>",
        f"<h1>{_esc(doc.title)}</h1>",
        _table_html(["Field", "Value"], _header_kv(doc)),
    ]
    fig_no = 0
    for b in doc.blocks:
        kind = b[0]
        if kind == "h":
            parts.append(f"<h{b[1] + 1}>{_esc(b[2])}</h{b[1] + 1}>")
        elif kind == "p":
            parts.append(f"<p>{_esc(b[1])}</p>")
        elif kind == "list":
            parts.append("<ul>" + "".join(f"<li>{_esc(i)}</li>" for i in b[1]) + "</ul>")
        elif kind == "kv":
            parts.append(_table_html(["Item", "Detail"], b[1]))
        elif kind == "table":
            parts.append(_table_html(b[1], b[2]))
            if b[3]:
                parts.append(f'<p class="cap">{_esc(b[3])}</p>')
        elif kind == "figure":
            fig_no += 1
            name = f"fig{fig_no}.png"
            archive.add((_png(b[1]), name))
            parts.append(f'<img src="{name}" width="420"/><p class="cap">{_esc(b[2])}</p>')
    story = pymupdf.Story("".join(parts), user_css=CSS, archive=archive)
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (50, 50, -50, -60)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    more = 1
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    pdf = pymupdf.open("pdf", buf.getvalue())
    for i, page in enumerate(pdf):
        page.insert_text(
            (50, mediabox.height - 30),
            f"{COMPANY} · {doc.id} Rev {doc.revision} · {doc.classification.upper()} · Page {i + 1} of {pdf.page_count}",
            fontsize=7.5,
            color=(0.35, 0.35, 0.35),
        )
    stamp = "D:20250101000000"
    pdf.set_metadata(
        {
            "title": doc.title,
            "author": doc.author,
            "subject": doc.id,
            "creator": COMPANY,
            "producer": "yokoten-datagen",
            "creationDate": stamp,
            "modDate": stamp,
        }
    )
    pdf.save(path, garbage=3, deflate=True, no_new_id=True)


# ---------------------------------------------------------------- DOCX
def _docx(doc: Doc, path: Path, rng: random.Random) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    d = Document()
    d.styles["Normal"].font.size = Pt(10)
    d.add_paragraph(f"{COMPANY} — {DOC_TYPE_LABEL[doc.doc_type]}").runs[0].bold = True
    d.add_heading(doc.title, level=0)
    _docx_table(d, ["Field", "Value"], _header_kv(doc))
    for b in doc.blocks:
        kind = b[0]
        if kind == "h":
            d.add_heading(b[2], level=b[1])
        elif kind == "p":
            d.add_paragraph(b[1])
        elif kind == "list":
            for i in b[1]:
                d.add_paragraph(i, style="List Bullet")
        elif kind == "kv":
            _docx_table(d, ["Item", "Detail"], b[1])
        elif kind == "table":
            _docx_table(d, b[1], b[2])
            if b[3]:
                d.add_paragraph(b[3]).runs[0].italic = True
        elif kind == "figure":
            d.add_picture(io.BytesIO(_png(b[1])), width=Inches(5.5))
            d.add_paragraph(b[2]).runs[0].italic = True
    cp = d.core_properties
    cp.title, cp.author, cp.subject = doc.title, doc.author, doc.id
    cp.created = cp.modified = FIXED_DATE
    d.save(path)


def _docx_table(d, header, rows) -> None:
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, header, strict=True):
        cell.text = str(h)
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        for cell, v in zip(t.add_row().cells, r, strict=True):
            cell.text = str(v)


# ---------------------------------------------------------------- XLSX
def _xlsx(doc: Doc, path: Path, rng: random.Random) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    info = wb.active
    info.title = "Document"
    for row in [("Company", COMPANY), ("Title", doc.title), *_header_kv(doc), ("Author", doc.author)]:
        info.append(row)
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 18, 60
    for name, header, rows in doc.sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DCE3EC")
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(header):
            width = max(len(str(h)), *(len(str(r[i])) for r in rows)) if rows else len(str(h))
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = min(max(width + 2, 8), 60)
        ws.freeze_panes = "A2"
    wb.properties.creator, wb.properties.title = doc.author, doc.title
    wb.properties.created = wb.properties.modified = FIXED_DATE
    wb.save(path)


# ---------------------------------------------------------------- Markdown
def _md(doc: Doc, path: Path, rng: random.Random) -> None:
    lines = [f"# {doc.title}", "", f"*{COMPANY} — {DOC_TYPE_LABEL[doc.doc_type]}*", ""]
    lines += _md_table(["Field", "Value"], _header_kv(doc)) + [""]
    for b in doc.blocks:
        kind = b[0]
        if kind == "h":
            lines += ["#" * (b[1] + 1) + " " + b[2], ""]
        elif kind == "p":
            lines += [b[1], ""]
        elif kind == "list":
            lines += [f"- {i}" for i in b[1]] + [""]
        elif kind == "kv":
            lines += _md_table(["Item", "Detail"], b[1]) + [""]
        elif kind == "table":
            lines += _md_table(b[1], b[2]) + ([f"*{b[3]}*"] if b[3] else []) + [""]
    path.write_text("\n".join(lines), "utf-8")


def _md_table(header, rows) -> list[str]:
    def cell(v):
        return str(v).replace("|", "/")

    return ["| " + " | ".join(map(cell, header)) + " |", "|" + "---|" * len(header)] + [
        "| " + " | ".join(map(cell, r)) + " |" for r in rows
    ]


# ---------------------------------------------------------------- scanned images
def _font(size: int, mono: bool = False, bold: bool = False) -> ImageFont.FreeTypeFont:
    from matplotlib import font_manager as fm

    name = "DejaVu Sans Mono" if mono else "DejaVu Sans"
    return ImageFont.truetype(
        fm.findfont(fm.FontProperties(family=name, weight="bold" if bold else "normal")), size
    )


def _scan(doc: Doc, path: Path, rng: random.Random) -> None:
    img = _draw_drawing(doc) if doc.scan["kind"] == "drawing" else _draw_inspection(doc)
    img = degrade(img, rng, doc.scan.get("skew", 3.0))
    if doc.fmt == "jpg":
        img.save(path, quality=62)
    else:
        img.save(path, optimize=True)


def degrade(img: Image.Image, rng: random.Random, skew: float) -> Image.Image:
    """Make a clean render look like a scan: skew, uneven light, blur, noise, speckles."""
    g = img.convert("L")
    angle = skew * (1 if rng.random() > 0.5 else -1)
    g = g.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=236)
    a = np.asarray(g).astype(np.float32)
    h, w = a.shape
    yy, xx = np.mgrid[0:h, 0:w]
    light = 1.0 - 0.16 * ((xx / w - rng.uniform(0.2, 0.8)) ** 2 + (yy / h - rng.uniform(0.2, 0.8)) ** 2)
    nrng = np.random.default_rng(rng.randrange(2**31))
    a = a * light - 14 + nrng.normal(0, 11, a.shape)
    speck = nrng.random(a.shape) < 0.0015
    a[speck] = nrng.uniform(20, 90, speck.sum())
    out = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))
    return out.filter(ImageFilter.GaussianBlur(0.7))


def _draw_drawing(doc: Doc) -> Image.Image:
    s = doc.scan
    W, H = 2400, 1650
    img = Image.new("L", (W, H), 250)
    d = ImageDraw.Draw(img)
    d.rectangle([30, 30, W - 30, H - 30], outline=0, width=4)
    # part view: plate with holes + dimensions
    x0, y0, x1, y1 = 300, 330, 1300, 900
    d.rectangle([x0, y0, x1, y1], outline=0, width=5)
    for cx in (420, 1180):
        for cy in (450, 780):
            d.ellipse([cx - 38, cy - 38, cx + 38, cy + 38], outline=0, width=4)
    d.ellipse([700, 520, 900, 720], outline=0, width=5)
    f = _font(34)
    d.line([x0, 980, x1, 980], fill=0, width=3)
    d.text(((x0 + x1) // 2 - 70, 990), s["dim_w"], font=f, fill=0)
    d.line([1380, y0, 1380, y1], fill=0, width=3)
    d.text((1395, (y0 + y1) // 2 - 20), s["dim_h"], font=f, fill=0)
    d.text((740, 740), s["dim_bore"], font=f, fill=0)
    # notes
    fn = _font(32)
    d.text((1500, 330), "NOTES:", font=_font(36, bold=True), fill=0)
    for i, n in enumerate(s["notes"]):
        d.text((1500, 390 + i * 52), f"{i + 1}. {n}", font=fn, fill=0)
    # title block
    tx, ty = 1300, 1130
    rows = [
        ("COMPANY", COMPANY.upper()),
        ("TITLE", s["title"].upper()),
        ("PART NO", s["part"]),
        ("MATERIAL", s["material"]),
        ("FINISH", s["finish"]),
        ("REV", s["rev"]),
        ("SCALE", s["scale"]),
        ("DRAWN", s["drawn"]),
        ("DATE", s["date"]),
    ]
    rh = 52
    d.rectangle([tx, ty, W - 60, ty + rh * len(rows)], outline=0, width=4)
    for i, (k, v) in enumerate(rows):
        yy = ty + i * rh
        d.line([tx, yy, W - 60, yy], fill=0, width=2)
        d.text((tx + 14, yy + 9), k, font=_font(28, bold=True), fill=0)
        d.text((tx + 250, yy + 7), v, font=_font(32, mono=True), fill=0)
    d.line([tx + 235, ty, tx + 235, ty + rh * len(rows)], fill=0, width=2)
    d.text((60, 60), f"DWG {s['part']}  SHEET 1/1", font=_font(34, bold=True), fill=0)
    return img


def _draw_inspection(doc: Doc) -> Image.Image:
    s = doc.scan
    W, H = 1700, 2200
    img = Image.new("L", (W, H), 252)
    d = ImageDraw.Draw(img)
    d.text((90, 80), COMPANY.upper(), font=_font(40, bold=True), fill=0)
    d.text((90, 140), "INCOMING INSPECTION RECORD", font=_font(46, bold=True), fill=0)
    y = 240
    for k, v in s["header"]:
        d.text((90, y), f"{k}:", font=_font(32, bold=True), fill=0)
        d.text((470, y), v, font=_font(32), fill=0)
        y += 54
    y += 30
    cols = [90, 620, 1030, 1340, W - 90]
    heads = ["CHARACTERISTIC", "SPECIFICATION", "MEASURED", "RESULT"]
    rh = 64
    d.rectangle([cols[0], y, cols[-1], y + rh * (len(s["rows"]) + 1)], outline=0, width=3)
    for i, hd in enumerate(heads):
        d.text((cols[i] + 12, y + 14), hd, font=_font(27, bold=True), fill=0)
    for r, row in enumerate(s["rows"]):
        yy = y + rh * (r + 1)
        d.line([cols[0], yy, cols[-1], yy], fill=0, width=2)
        for i, v in enumerate(row):
            d.text((cols[i] + 12, yy + 14), str(v), font=_font(29, mono=i > 0), fill=0)
    for c in cols[1:-1]:
        d.line([c, y, c, y + rh * (len(s["rows"]) + 1)], fill=0, width=2)
    y += rh * (len(s["rows"]) + 1) + 70
    d.text((90, y), f"DECISION: {s['decision']}", font=_font(40, bold=True), fill=0)
    d.rectangle([80, y - 12, 80 + 60 + 26 * len(f"DECISION: {s['decision']}"), y + 60], outline=0, width=4)
    d.text((90, y + 130), f"Inspector: {s['inspector']}", font=_font(32), fill=0)
    d.text((90, y + 190), f"Remarks: {s['remarks']}", font=_font(30), fill=0)
    # a faint hand-drawn tick/signature
    pts = [(900 + 12 * i, y + 170 + int(18 * math.sin(i / 2))) for i in range(30)]
    d.line(pts, fill=40, width=4)
    return img


# ---------------------------------------------------------------- matplotlib figures
def figure_png(draw, size=(6.0, 3.2)) -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=size, dpi=110)
    draw(ax)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()
