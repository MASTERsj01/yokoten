"""Format loaders -> ordered layout elements, exposed as a LangChain BaseLoader.

Every element keeps its location (PDF page + bbox in points, image bbox normalised 0..1, sheet row, section path)
so citations can open the exact place in the source viewer.
"""

import re
import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path

from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document as LCDocument

from yokoten.config import settings
from yokoten.ingest import ocr as ocrmod

SUPPORTED = {".pdf", ".docx", ".xlsx", ".csv", ".md", ".txt", ".png", ".jpg", ".jpeg"}


@dataclass
class Element:
    text: str
    kind: str = "text"  # heading | text | list | table | row | figure | ocr
    page: int | None = None
    bbox: tuple | None = None
    section: str = ""
    level: int = 0
    extra: dict = field(default_factory=dict)


@dataclass
class Parsed:
    elements: list[Element]
    n_pages: int | None = None
    figures: list[dict] = field(default_factory=list)  # {"page", "bbox", "caption", "path"}
    ocr: dict | None = None  # engine, confidence, angle, steps, fields, text


def clean(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)  # also splits ligatures (ﬁ -> fi)
    return re.sub(r"[ \t]+", " ", text).strip()


class _Sections:
    """Tracks the heading path while walking a document."""

    def __init__(self):
        self.stack: list[tuple[int, str]] = []

    def push(self, level: int, text: str) -> None:
        self.stack = [(lv, t) for lv, t in self.stack if lv < level] + [(level, text)]

    @property
    def path(self) -> str:
        """Heading path below the document title (level 1), e.g. 'D4 - Root cause > Detection'."""
        inner = [t for lv, t in self.stack if lv > 1]
        return " > ".join(inner) if inner else " > ".join(t for lv, t in self.stack if lv > 0)


def _table_text(rows: list[list]) -> str:
    rows = [[clean(str(c or "")) for c in r] for r in rows if any(c not in (None, "") for c in r)]
    return "\n".join(" | ".join(r) for r in rows)


# ------------------------------------------------------------------ PDF
def parse_pdf(path: Path, fig_dir: Path) -> Parsed:
    import pdfplumber
    import pymupdf

    elements: list[Element] = []
    figures: list[dict] = []
    sec = _Sections()
    doc = pymupdf.open(path)
    with pdfplumber.open(path) as pl:
        for pno, page in enumerate(doc, start=1):
            height = page.rect.height
            tables = [
                (t.bbox, rows)
                for t in pl.pages[pno - 1].find_tables()
                if len(rows := [r for r in t.extract() if any(r)]) >= 2
                and max(sum(1 for c in r if c) for r in rows) >= 2
            ]
            items: list[tuple[float, Element]] = []
            for bbox, rows in tables:
                items.append((bbox[1], Element(_table_text(rows), "table", pno, tuple(bbox))))

            def in_table(b, tables=tables):
                cx, cy = (b[0] + b[2]) / 2, (b[1] + b[3]) / 2
                return any(t[0] <= cx <= t[2] and t[1] <= cy <= t[3] for t, _ in tables)

            for block in page.get_text("dict")["blocks"]:
                b = tuple(block["bbox"])
                if block["type"] == 1:  # image -> figure
                    if b[3] - b[1] < 40:
                        continue
                    fig_dir.mkdir(parents=True, exist_ok=True)
                    fpath = fig_dir / f"p{pno}_{len(figures) + 1}.{block.get('ext', 'png')}"
                    fpath.write_bytes(block["image"])
                    el = Element("", "figure", pno, b, extra={"path": fpath})
                    items.append((b[1], el))
                    continue
                if b[1] > height - 45 or in_table(b):  # footer line / text already captured in a table
                    continue
                spans = [s for ln in block["lines"] for s in ln["spans"]]
                text = clean(" ".join("".join(s["text"] for s in ln["spans"]) for ln in block["lines"]))
                if not text:
                    continue
                size = max(s["size"] for s in spans)
                level = 1 if size >= 15 else 2 if size >= 12 else 3 if size >= 10.8 else 0
                items.append((b[1], Element(text, "heading" if level else "text", pno, b, level=level)))
            items.sort(key=lambda x: x[0])
            for _, el in items:
                if el.kind == "heading":
                    sec.push(el.level, el.text)
                el.section = sec.path
                if (
                    el.kind == "text"
                    and el.text.lower().startswith("figure")
                    and elements
                    and elements[-1].kind == "figure"
                    and not elements[-1].text
                ):
                    elements[-1].text = el.text  # caption belongs to the figure above it
                    continue
                elements.append(el)
    for el in elements:
        if el.kind == "figure":
            el.text = el.text or f"Figure on page {el.page}"
            figures.append(
                {"page": el.page, "bbox": list(el.bbox), "caption": el.text, "path": el.extra.pop("path")}
            )
    return Parsed(elements, n_pages=doc.page_count, figures=figures)


# ------------------------------------------------------------------ DOCX
def parse_docx(path: Path) -> Parsed:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(str(path))
    elements: list[Element] = []
    sec = _Sections()
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = clean(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style == "Title" or style.startswith("Heading"):
                level = (
                    1
                    if style == "Title"
                    else min(int(style.split()[-1]) + 1, 4)
                    if style[-1].isdigit()
                    else 2
                )
                sec.push(level, text)
                elements.append(Element(text, "heading", section=sec.path, level=level))
            else:
                kind = "list" if "List" in style else "text"
                elements.append(Element(("- " if kind == "list" else "") + text, kind, section=sec.path))
        elif tag == "tbl":
            t = Table(child, d)
            rows = [[c.text for c in r.cells] for r in t.rows]
            elements.append(Element(_table_text(rows), "table", section=sec.path))
    return Parsed(elements)


# ------------------------------------------------------------------ XLSX / CSV
def _rows_to_elements(sheet: str, header: list, rows: list[list]) -> list[Element]:
    out = []
    for i, r in enumerate(rows, start=2):
        pairs = [(str(h), v) for h, v in zip(header, r, strict=False) if v not in (None, "")]
        if not pairs:
            continue
        text = "; ".join(f"{h}: {v}" for h, v in pairs)
        out.append(
            Element(text, "row", section=sheet, extra={"row": i, "sheet": sheet, "values": dict(pairs)})
        )
    return out


def parse_xlsx(path: Path) -> Parsed:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    elements: list[Element] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        if ws.title == "Document" or len(rows[0]) == 2:
            elements.append(Element(_table_text(rows), "table", section=ws.title))
        else:
            elements.append(
                Element(
                    f"Sheet {ws.title}: columns " + ", ".join(map(str, rows[0])),
                    "heading",
                    section=ws.title,
                    level=2,
                )
            )
            elements += _rows_to_elements(ws.title, rows[0], rows[1:])
    return Parsed(elements)


def parse_csv(path: Path) -> Parsed:
    import pandas as pd

    df = pd.read_csv(path).fillna("")
    return Parsed(_rows_to_elements(path.stem, list(df.columns), df.values.tolist()))


# ------------------------------------------------------------------ Markdown / text
def parse_md(path: Path) -> Parsed:
    elements: list[Element] = []
    sec = _Sections()
    para: list[str] = []
    table: list[str] = []

    def flush():
        if para:
            elements.append(Element(clean(" ".join(para)), "text", section=sec.path))
            para.clear()
        if table:
            rows = [
                [c.strip() for c in ln.strip().strip("|").split("|")]
                for ln in table
                if not re.match(r"^\|?[-\s|]+\|?$", ln)
            ]
            elements.append(Element(_table_text(rows), "table", section=sec.path))
            table.clear()

    for line in path.read_text("utf-8", errors="replace").splitlines():
        s = line.strip()
        if m := re.match(r"^(#{1,6})\s+(.*)", s):
            flush()
            level = len(m.group(1))
            sec.push(level, clean(m.group(2)))
            elements.append(Element(clean(m.group(2)), "heading", section=sec.path, level=level))
        elif s.startswith("|"):
            if para:
                flush()
            table.append(s)
        elif s.startswith(("- ", "* ")):
            flush()
            elements.append(Element(clean(s), "list", section=sec.path))
        elif not s:
            flush()
        else:
            if table:
                flush()
            para.append(s.strip("*"))
    flush()
    return Parsed(elements)


# ------------------------------------------------------------------ images (OCR)
def parse_image(path: Path, steps_dir: Path, engine: str = "auto") -> Parsed:
    res = ocrmod.ocr(path, engine=engine)
    fields = ocrmod.extract_fields(res.text)
    elements = [
        Element(ln.text, "ocr", page=1, bbox=ln.bbox, section="OCR", extra={"conf": ln.conf})
        for ln in res.lines
    ]
    step_paths = ocrmod.save_steps(res, steps_dir)
    return Parsed(
        elements,
        n_pages=1,
        ocr={
            "engine": res.engine,
            "confidence": round(res.confidence, 3),
            "angle": res.angle,
            "steps": step_paths,
            "fields": fields,
            "text": res.text,
        },
    )


def parse(path: Path, work_dir: Path, ocr_engine: str = "auto") -> Parsed:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path, work_dir / "figures")
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".xlsx":
        return parse_xlsx(path)
    if ext == ".csv":
        return parse_csv(path)
    if ext in (".md", ".txt"):
        return parse_md(path)
    if ext in (".png", ".jpg", ".jpeg"):
        return parse_image(path, work_dir / "ocr", ocr_engine)
    raise ValueError(f"Unsupported file type: {ext}")


class EngineeringDocLoader(BaseLoader):
    """LangChain loader: one LangChain Document per layout element, with location metadata."""

    def __init__(self, path: Path, doc_meta: dict | None = None, ocr_engine: str = "auto"):
        self.path = Path(path)
        self.doc_meta = doc_meta or {}
        self.ocr_engine = ocr_engine
        self.parsed: Parsed | None = None

    def lazy_load(self) -> Iterator[LCDocument]:
        key = self.doc_meta.get("rev_key", self.path.stem)
        self.parsed = parse(self.path, settings.var_dir / "work" / key.replace("@", "_"), self.ocr_engine)
        for i, el in enumerate(self.parsed.elements):
            yield LCDocument(
                page_content=el.text,
                metadata={
                    **self.doc_meta,
                    "element": i,
                    "kind": el.kind,
                    "page": el.page,
                    "bbox": list(el.bbox) if el.bbox else None,
                    "section": el.section,
                    "level": el.level,
                    **({"row": el.extra["row"], "values": el.extra["values"]} if "row" in el.extra else {}),
                },
            )
